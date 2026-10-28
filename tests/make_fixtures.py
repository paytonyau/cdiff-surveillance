"""Regenerate the tiny spreadsheet fixtures used by tests/testthat/test-parse.R.

    python3 tests/make_fixtures.py

Two layouts are mimicked, taken from the real files:
  new_layout.xlsx : 2024+ style  (sheets Cover | Table_1_national_data | Table_2_raw_data, header on row 4,
                                  hyphenated metric labels, "National" rows, UKHSA centre / ICB / Region name)
  old_layout.xlsx : 2019-2021 style (sheets caveats | monthly_table | data, header on row 1, "PHE Centre" + "STP",
                                  no ICB / Region name, "CCG" rows, unhyphenated metric labels)
"""
import openpyxl, pathlib
out = pathlib.Path(__file__).parent / "testthat" / "fixtures"; out.mkdir(parents=True, exist_ok=True)

COMP = ["hoha", "coha", "coia", "coca", "unknown", "noinfo"]
def counts(seed):                      # deterministic components; total is their sum
    base = [5 + seed, 2, 1, 3 + seed % 2, 0, 0]
    return dict(zip(COMP, base)), sum(base)

LABELS_NEW = {"total": "Total cases", "hoha": "Hospital-onset, healthcare-associated", "coha": "Community-onset, healthcare-associated",
              "coia": "Community-onset, indeterminate-association", "coca": "Community-onset, community-associated",
              "unknown": "Unknown", "noinfo": "No information"}
LABELS_OLD = {"total": "Total cases", "hoha": "Hospital-onset, healthcare associated", "coha": "Community-onset, healthcare associated",
              "coia": "Community-onset, indeterminate association", "coca": "Community-onset, community associated",
              "unknown": "Unknown", "noinfo": "No information"}

def rows(labels, trust_type, geo_type, new):
    out = []
    for month in (4, 5):
        for code, name, kind in (("R1A", "  Alpha  NHS Foundation Trust ", trust_type), ("R1B", "Beta Hospitals NHS Trust", trust_type),
                                 ("00A", "Alpha sub ICB", geo_type)):
            if kind == geo_type:      # the sub-ICB holds the same cases as its trusts: counts must equal their sum
                parts = [counts(ord(c) + month) for c in "AB"]
                comp = {k: sum(p[0][k] for p in parts) for k in COMP}; total = sum(p[1] for p in parts)
            else:
                comp, total = counts(ord(code[-1]) + month)
            for key, label in labels.items():
                fig = total if key == "total" else comp[key]
                base = ["C. difficile", kind, code, name]
                if new:
                    out.append(base + ["North West", "NHS ICB", "North West", 2024, month, label, fig, "Yes" if kind != geo_type else None, "No"])
                else:
                    out.append(base + ["Cheshire and Merseyside", "STP1", 2019, month, label, fig, "Yes", "No"])
    if new:   # National rows (2024+ files only)
        for month in (4, 5):
            tot = sum(counts(ord(c) + month)[1] for c in "AB")
            for key, label in labels.items():
                fig = sum((counts(ord(c) + month)[1] if key == "total" else counts(ord(c) + month)[0][key]) for c in "AB")
                out.append(["C. difficile", "National", "ENG", "England", None, None, None, 2024, month, label, fig, None, "No"])
    return out

# ---- new layout
wb = openpyxl.Workbook(); wb.active.title = "Cover"; wb["Cover"].append(["UKHSA monthly counts"])
w1 = wb.create_sheet("Table_1_national_data"); w1.append(["National table"]); w1.append([]); w1.append([]); w1.append(["Organisation name", "Organisation code", "Metric", "April 2024"])
w2 = wb.create_sheet("Table_2_raw_data"); [w2.append([t]) for t in ("Raw data", "Subtitle", "One table")]
w2.append(["Collection", "Organisation type", "Organisation code", "Organisation name", "UKHSA centre", "ICB", "Region name", "Year", "Month", "Metric", "Figure", "Data signed off", "Data changed since last publication?"])
for r in rows(LABELS_NEW, "NHS acute trust", "Sub-ICB", True): w2.append(r)
wb.save(out / "new_layout.xlsx")

# ---- old layout
wb = openpyxl.Workbook(); wb.active.title = "caveats"; wb["caveats"].append(["caveat text"])
wb.create_sheet("monthly_table").append(["pivot, not used"])
w = wb.create_sheet("data")
w.append(["Collection", "Organisation type", "Organisation code", "Organisation name", "PHE Centre", "STP", "Year", "Month", "Metric", "Figure", "Data signed off", "Data changed since last publication?"])
for r in rows(LABELS_OLD, "NHS acute trust", "CCG", False): w.append(r)
wb.save(out / "old_layout.xlsx")

# ---- broken variants (must fail loudly)
wb = openpyxl.load_workbook(out / "new_layout.xlsx"); ws = wb["Table_2_raw_data"]
for row in ws.iter_rows(min_row=5):
    if row[9].value == "Unknown": row[9].value = "Mystery category"
wb.save(out / "bad_metric.xlsx")
wb = openpyxl.load_workbook(out / "new_layout.xlsx"); ws = wb["Table_2_raw_data"]
ws.cell(row=4, column=11).value = "Count"          # 'Figure' column renamed -> no header row found
wb.save(out / "no_header.xlsx")
print("fixtures written to", out)
